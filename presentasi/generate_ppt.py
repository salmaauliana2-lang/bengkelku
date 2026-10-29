"""Membuat deck presentasi WorkshopPro (Tugas Akhir BLKPP).

Jalankan:
    pip install python-pptx
    python presentasi/generate_ppt.py
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- palet
NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x2E, 0x90, 0xFA)
AMBER = RGBColor(0xF7, 0x90, 0x09)
GREEN = RGBColor(0x12, 0xB7, 0x6A)
RED = RGBColor(0xF0, 0x44, 0x38)
VIOLET = RGBColor(0x7A, 0x5A, 0xF8)
LINE = RGBColor(0xDD, 0xE4, 0xEE)
FONT = "Segoe UI"

SW, SH = 13.333, 7.5
ML = 0.7  # margin kiri


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
         line_w=1.0):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def text(slide, x, y, w, h, blocks, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP):
    """blocks: list of dict(t=, size=, bold=, color=, space_before=, space_after=)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, b in enumerate(blocks):
        # baris baru ("\\n") dipecah menjadi paragraf terpisah
        for j, line_txt in enumerate(b["t"].split("\n")):
            p = tf.paragraphs[0] if (i == 0 and j == 0) else tf.add_paragraph()
            p.alignment = b.get("align", align)
            p.space_before = Pt(b.get("space_before", 0) if j == 0 else 0)
            p.space_after = Pt(b.get("space_after", 0) if j == 0 else 3)
            p.line_spacing = b.get("line_spacing", 1.12)
            run = p.add_run()
            run.text = line_txt
            run.font.name = FONT
            run.font.size = Pt(b.get("size", 14))
            run.font.bold = b.get("bold", False)
            run.font.color.rgb = b.get("color", NAVY)
    return tb


def header(slide, kicker, title, accent=BLUE):
    rect(slide, 0, 0, SW, 1.34, fill=WHITE)
    rect(slide, 0, 1.32, SW, 0.03, fill=LINE)
    rect(slide, ML, 0.42, 0.06, 0.62, fill=accent)
    text(slide, ML + 0.22, 0.40, 10.5, 0.3,
         [{"t": kicker.upper(), "size": 11, "bold": True, "color": accent}])
    text(slide, ML + 0.22, 0.66, 11.6, 0.5,
         [{"t": title, "size": 26, "bold": True, "color": NAVY}])


def footer(slide, page):
    rect(slide, 0, 6.96, SW, 0.03, fill=LINE)
    text(slide, ML, 7.06, 8.5, 0.3,
         [{"t": "WorkshopPro  ·  Tugas Akhir Pelatihan BLKPP  ·  "
                "Salma Auliana", "size": 9.5, "color": SLATE}])
    text(slide, 11.6, 7.06, 1.0, 0.3,
         [{"t": f"{page:02d}", "size": 9.5, "bold": True, "color": SLATE}],
         align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, accent, title, body, title_size=14.5,
         body_size=11.5, tag=None):
    rect(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y, 0.07, h, fill=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x + 0.26, y + 0.2, w - 0.5, 0.35,
         [{"t": title, "size": title_size, "bold": True, "color": NAVY}])
    if tag:
        text(slide, x + 0.26, y + 0.62, w - 0.5, 0.25,
             [{"t": tag, "size": 9.5, "bold": True, "color": accent}])
        body_y = y + 0.92
    else:
        body_y = y + 0.66
    blocks = []
    for i, line_txt in enumerate(body):
        blocks.append({
            "t": f"•  {line_txt}",
            "size": body_size,
            "color": SLATE,
            "space_after": 5,
            "line_spacing": 1.1,
        })
        _ = i
    text(slide, x + 0.26, body_y, w - 0.5, h - (body_y - y) - 0.15, blocks)


def bullets(slide, x, y, w, h, items, size=13.5, color=SLATE, gap=9,
            marker="•", bold_head=True):
    """items: list of str, atau tuple (heading, isi) — heading dibuat tebal."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.14

        parts = []
        if isinstance(it, tuple) and bold_head:
            parts.append((f"{marker}  {it[0]}", True, NAVY))
            parts.append((f" — {it[1]}", False, color))
        else:
            body = f"{marker}  {it[0]} — {it[1]}" if isinstance(it, tuple) \
                else f"{marker}  {it}"
            parts.append((body, False, color))

        for txt, bold, col in parts:
            r = p.add_run()
            r.text = txt
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = col
    return tb


def make_table(slide, x, y, w, h, data, col_w, header_fill=NAVY,
               size=11, header_size=11, row_h=0.36, align_first_left=True):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    tbl = shape.table
    tbl.first_row = True
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(rows):
        tbl.rows[r].height = Inches(row_h if r else row_h + 0.04)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.0
                if c > 0 or not align_first_left:
                    p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(header_size if r == 0 else size)
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else NAVY
    return tbl


# ================================================================ deck
prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)

# ------------------------------------------------------------ 1. judul
s = blank(prs)
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, 0, 0.22, SH, fill=BLUE)
rect(s, 8.9, -1.4, 5.6, 5.6, fill=RGBColor(0x16, 0x21, 0x3A),
     shape=MSO_SHAPE.OVAL)
rect(s, 10.6, 3.4, 3.8, 3.8, fill=RGBColor(0x1B, 0x2A, 0x47),
     shape=MSO_SHAPE.OVAL)

text(s, 1.0, 1.5, 9.5, 0.4,
     [{"t": "TUGAS AKHIR PELATIHAN BLKPP", "size": 13, "bold": True,
       "color": AMBER}])
text(s, 1.0, 2.0, 10.6, 1.9, [
    {"t": "WorkshopPro", "size": 54, "bold": True, "color": WHITE,
     "space_after": 4},
    {"t": "Sistem Manajemen Bengkel Motor", "size": 30, "bold": True,
     "color": BLUE},
])
text(s, 1.0, 4.15, 9.4, 1.0, [
    {"t": "Aplikasi web full-stack untuk mengelola inventaris spare part, "
          "penjualan, pekerjaan servis, dan pembelian — dengan stok "
          "terlacak real-time serta laporan otomatis.",
     "size": 15, "color": RGBColor(0xC3, 0xCE, 0xE0),
     "line_spacing": 1.35},
])

chips = [("Go 1.27 · Gin · PostgreSQL", BLUE),
         ("React 19 · TypeScript · Vite", GREEN),
         ("JWT · RBAC · Docker", AMBER)]
cx = 1.0
for label, col in chips:
    w = 0.16 * len(label) * 0.5 + 0.55
    rect(s, cx, 5.35, w, 0.44, fill=RGBColor(0x16, 0x21, 0x3A), line=col,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, cx, 5.44, w, 0.3, [{"t": label, "size": 11.5, "bold": True,
                                "color": WHITE}], align=PP_ALIGN.CENTER)
    cx += w + 0.25

text(s, 1.0, 6.35, 8.0, 0.6, [
    {"t": "Salma Auliana  ·  salma.auliana2@gmail.com  ·  Tugas Akhir BLKPP",
     "size": 12, "bold": True, "color": WHITE, "space_after": 3},
    {"t": "Akun demo:  owner@bengkelku.id / owner123", "size": 11,
     "color": RGBColor(0x9A, 0xAB, 0xC6)},
])

# -------------------------------------------------- 2. latar belakang
s = blank(prs)
header(s, "Latar belakang", "Masalah yang Dihadapi Bengkel", AMBER)
text(s, ML, 1.65, 5.4, 0.4,
     [{"t": "Kondisi sebelum ada sistem", "size": 14, "bold": True,
       "color": AMBER}])
bullets(s, ML, 2.1, 5.5, 4.2, [
    ("Pencatatan manual", "nota & buku kas sering hilang atau salah tulis"),
    ("Stok tidak akurat", "spare part mudah habis atau menumpuk diam di gudang"),
    ("Servis tidak terlacak", "pelanggan sulit mengetahui progres kendaraannya"),
    ("Laporan lambat", "rekap penjualan dihitung manual tiap akhir bulan"),
    ("Akses tidak terkontrol", "semua orang bisa mengubah data yang sama"),
], size=13, gap=13)

rect(s, 6.6, 1.65, 6.05, 4.8, fill=LIGHT, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 6.9, 1.9, 5.5, 0.4,
     [{"t": "Dampaknya", "size": 14, "bold": True, "color": NAVY}])
impacts = [
    ("Selisih stok", "Barang terjual tidak masuk catatan.", RED),
    ("Kerugian", "Harga jual & modal tidak terkontrol.", AMBER),
    ("Pelayanan buruk", "Antrian servis tidak tertib.", VIOLET),
    ("Keputusan lambat", "Pemilik tidak punya data real-time.", BLUE),
]
iy = 2.4
for title_i, desc, col in impacts:
    rect(s, 6.9, iy, 5.45, 0.86, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, 6.9, iy, 0.07, 0.86, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 7.15, iy + 0.14, 5.0, 0.3,
         [{"t": title_i, "size": 13, "bold": True, "color": NAVY}])
    text(s, 7.15, iy + 0.46, 5.0, 0.3,
         [{"t": desc, "size": 11.5, "color": SLATE}])
    iy += 1.02

footer(s, 2)

# ------------------------------------------------------- 3. tujuan
s = blank(prs)
header(s, "Tujuan", "Solusi & Sasaran Proyek", GREEN)
goals = [
    ("1", "Digitalisasi operasional", "Satu aplikasi untuk seluruh aktivitas "
     "bengkel: penjualan, servis, pembelian, dan stok.", BLUE),
    ("2", "Stok akurat & terlacak", "Setiap pergerakan barang tercatat "
     "otomatis lengkap dengan stok sebelum–sesudah.", GREEN),
    ("3", "Visibilitas pekerjaan servis", "Status antrian kendaraan transparan "
     "bagi admin, mekanik, maupun pelanggan.", AMBER),
    ("4", "Keputusan berbasis data", "Dashboard ringkasan + tren penjualan "
     "7 hari terakhir secara otomatis.", VIOLET),
]
gx = ML
for num, title_i, desc, col in goals:
    rect(s, gx, 1.7, 2.9, 2.5, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, gx + 0.25, 1.95, 0.62, 0.62, fill=col, shape=MSO_SHAPE.OVAL)
    text(s, gx + 0.25, 2.07, 0.62, 0.4,
         [{"t": num, "size": 17, "bold": True, "color": WHITE}],
         align=PP_ALIGN.CENTER)
    text(s, gx + 0.25, 2.75, 2.4, 0.5,
         [{"t": title_i, "size": 14, "bold": True, "color": NAVY,
           "line_spacing": 1.1}])
    text(s, gx + 0.25, 3.3, 2.42, 0.85,
         [{"t": desc, "size": 11, "color": SLATE, "line_spacing": 1.2}])
    gx += 3.03

rect(s, ML, 4.5, 11.93, 2.15, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, ML + 0.4, 4.75, 11.1, 0.4,
     [{"t": "Cakupan aplikasi (6 modul)", "size": 14, "bold": True,
       "color": AMBER}])
mods = ["Dashboard", "Inventaris", "Penjualan", "Servis Bengkel",
        "Pembelian", "Pengguna & Role"]
mx = ML + 0.4
mw = 1.78
for i, m in enumerate(mods):
    rect(s, mx, 5.3, mw, 0.72, fill=RGBColor(0x1B, 0x2A, 0x47), line=BLUE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, mx, 5.5, mw, 0.35,
         [{"t": m, "size": 12, "bold": True, "color": WHITE}],
         align=PP_ALIGN.CENTER)
    mx += mw + 0.08
text(s, ML + 0.4, 6.16, 11.1, 0.3,
     [{"t": "Berjalan penuh di satu aplikasi web — tanpa pencatatan manual.",
       "size": 11.5, "color": RGBColor(0xB9, 0xC7, 0xDC)}])
footer(s, 3)

# ------------------------------------------------- 4. fitur utama
s = blank(prs)
header(s, "Fitur", "Enam Modul Utama Aplikasi", BLUE)
features = [
    ("Dashboard", BLUE, ["4 kartu statistik (penjualan hari ini, servis "
     "aktif, stok menipis, pembelian bulan ini)",
     "Grafik tren penjualan & kunjungan 7 hari",
     "Donut status servis + transaksi terbaru"]),
    ("Inventaris", GREEN, ["56 jenis spare part: SKU, kategori, harga beli "
     "dan jual",
     "Riwayat pergerakan stok (masuk/keluar/adjustment)",
     "Peringatan stok menipis & habis"]),
    ("Penjualan", AMBER, ["Kasir cepat: pilih item, otomatis kalkulasi subtotal",
     "Invoice INV- otomatis, diskon, catatan",
     "Stok terpotong + tercatat dalam satu transaksi DB"]),
    ("Servis Bengkel", VIOLET, ["Antrian kendaraan: menunggu → dikerjakan → "
     "selesai/batal",
     "Keluhan, diagnosa, jasa kerja & spare part terpakai",
     "Pembayaran tunai + status lunas/belum"]),
    ("Pembelian", RED, ["Restock dari supplier beserta rincian item",
     "Stok otomatis bertambah tercatat sebagai 'masuk'",
     "Riwayat pembelian per tanggal & supplier"]),
    ("Pengguna & Role", NAVY, ["4 role: owner, admin, mekanik, staff",
     "Owner mengelola role lewat menu Pengguna",
     "Menu menyesuaikan role yang sedang login"]),
]
positions = [(ML, 1.62), (4.74, 1.62), (8.78, 1.62),
             (ML, 4.3), (4.74, 4.3), (8.78, 4.3)]
for (title_i, col, body), (x, y) in zip(features, positions):
    card(s, x, y, 3.84, 2.5, col, title_i, body, body_size=10.8)
footer(s, 4)

# --------------------------------------------- 5. alur operasional
s = blank(prs)
header(s, "Alur bisnis", "Alur Operasional Bengkel", AMBER)
steps = [
    ("Pembelian", "Supplier → barang masuk", GREEN),
    ("Inventaris", "Stok & harga diperbarui", BLUE),
    ("Penjualan", "Transaksi kasir spare part", AMBER),
    ("Servis", "Kendaraan diperiksa & dikerjakan", VIOLET),
    ("Pembayaran", "Pelunasan tunai tercatat", RED),
    ("Dashboard", "Ringkasan & tren otomatis", NAVY),
]
sx = ML
sw = 1.92
descs_short = ["Barang masuk dari supplier", "Stok & harga diperbarui",
               "Kasir spare part", "Kendaraan dikerjakan",
               "Pelunasan tercatat", "Ringkasan & tren otomatis"]
for i, (title_i, desc, col) in enumerate(steps):
    shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    sh = rect(s, sx, 1.8, sw + 0.16, 1.15, fill=col, shape=shape)
    try:
        #kecilkan kedalaman ujung panah agar ruang teks lebih lebar
        sh.adjustments[0] = 0.18
    except (IndexError, AttributeError):
        pass
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = title_i
    r.font.name = FONT
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.line_spacing = 1.05
    r2 = p2.add_run()
    r2.text = descs_short[i]
    r2.font.name = FONT
    r2.font.size = Pt(9)
    r2.font.color.rgb = WHITE
    _ = desc
    sx += sw

notes = [
    ("Transaksi atomik", "Penjualan, pemotongan stok, dan riwayat stok "
     "berjalan dalam SATU transaksi PostgreSQL — bila gagal, semuanya "
     "kembali seperti semula."),
    ("Kunci baris (FOR UPDATE)", "Baris produk dikunci saat proses agar dua "
     "kasir tidak bisa menjual stok yang sama di saat bersamaan."),
    ("Audit trail", "Setiap pergerakan stok menyimpan siapa pelaku, "
     "stok sebelum, dan stok sesudah."),
]
ny = 3.4
for title_i, desc in notes:
    rect(s, ML, ny, 11.93, 1.0, fill=LIGHT, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, ML, ny, 0.07, 1.0, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, ML + 0.3, ny + 0.17, 3.0, 0.3,
         [{"t": title_i, "size": 13.5, "bold": True, "color": NAVY}])
    text(s, ML + 3.4, ny + 0.16, 8.2, 0.7,
         [{"t": desc, "size": 12, "color": SLATE, "line_spacing": 1.2}])
    ny += 1.12
footer(s, 5)

# --------------------------------------------------- 6. arsitektur
s = blank(prs)
header(s, "Arsitektur", "Arsitektur Sistem", BLUE)

rect(s, ML, 2.0, 3.0, 1.7, fill=WHITE, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, ML + 0.2, 2.18, 2.6, 0.3,
     [{"t": "Browser (Klien)", "size": 13, "bold": True, "color": BLUE}],
     align=PP_ALIGN.CENTER)
text(s, ML + 0.2, 2.55, 2.6, 1.0,
     [{"t": "SPA React 19 + TypeScript\nReact Router (halaman terproteksi)\n"
            "State sesi & token JWT", "size": 11, "color": SLATE,
       "line_spacing": 1.3}], align=PP_ALIGN.CENTER)

arrow1 = rect(s, ML + 3.1, 2.6, 0.9, 0.5, fill=BLUE, shape=MSO_SHAPE.RIGHT_ARROW)

rect(s, ML + 4.1, 2.0, 3.4, 1.7, fill=WHITE, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, ML + 4.3, 2.18, 3.0, 0.3,
     [{"t": "REST API — Go + Gin", "size": 13, "bold": True, "color": GREEN}],
     align=PP_ALIGN.CENTER)
text(s, ML + 4.3, 2.55, 3.0, 1.0,
     [{"t": "Middleware: CORS → JWT → RBAC\nHandler → Repository (pgx)\n"
            "Validasi input binding", "size": 11, "color": SLATE,
       "line_spacing": 1.3}], align=PP_ALIGN.CENTER)

arrow2 = rect(s, ML + 7.6, 2.6, 0.9, 0.5, fill=BLUE,
              shape=MSO_SHAPE.RIGHT_ARROW)

rect(s, ML + 8.6, 2.0, 3.33, 1.7, fill=WHITE, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, ML + 8.8, 2.18, 2.9, 0.3,
     [{"t": "PostgreSQL", "size": 13, "bold": True, "color": AMBER}],
     align=PP_ALIGN.CENTER)
text(s, ML + 8.8, 2.55, 2.9, 1.0,
     [{"t": "10 tabel + index & CHECK\nMigrasi SQL otomatis saat start\n"
            "Koneksi pool pgx v5", "size": 11, "color": SLATE,
       "line_spacing": 1.3}], align=PP_ALIGN.CENTER)

text(s, ML, 4.05, 6.0, 0.35,
     [{"t": "Prinsip yang diterapkan", "size": 14, "bold": True,
       "color": NAVY}])
bullets(s, ML, 4.5, 5.7, 2.3, [
    "Terpisah tugas: handler (HTTP) · repository (SQL) · model (data)",
    "Semua route privat dijaga JWT, lalu difilter per role",
    "Migrasi berjalan otomatis → instalasi baru langsung siap",
    "Graceful shutdown + timeout server untuk keamanan proses",
], size=12.5, gap=11)

rect(s, 6.8, 4.05, 5.83, 2.6, fill=LIGHT, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.1, 4.25, 5.3, 0.35,
     [{"t": "Detail teknis penting", "size": 14, "bold": True,
       "color": NAVY}])
bullets(s, 7.1, 4.7, 5.2, 1.8, [
    "Token JWT HS256, berlaku 8 jam, wajib Bearer",
    "Bcrypt untuk semua password",
    "Endpoint /health mengecek koneksi database",
    "Mode satu wadah: backend menyajikan hasil build frontend (STATIC_DIR)",
], size=12, gap=10)
footer(s, 6)

# -------------------------------------------------- 7. teknologi
s = blank(prs)
header(s, "Teknologi", "Stack yang Digunakan", VIOLET)
stacks = [
    ("Backend (Go)", GREEN, [
        "Go 1.27 + Gin (router & middleware)",
        "pgx v5 + connection pool PostgreSQL",
        "golang-jwt/v5 (JWT HS256)",
        "golang.org/x/crypto (bcrypt)",
        "gin-contrib/cors, godotenv",
        "26 endpoint API, dibagi per paket modul",
    ]),
    ("Frontend (React)", BLUE, [
        "React 19 + TypeScript 6",
        "Vite 8 (build super cepat)",
        "React Router 7 (rute terproteksi)",
        "Komponen: chart SVG, sparkline, badge",
        "CSS modern responsif (sidebar mobile)",
        "ESLint untuk kualitas kode",
    ]),
    ("DevOps & Deployment", AMBER, [
        "Dockerfile multi-stage (frontend → Go → Alpine)",
        "Satu image ringan, non-root user",
        "Railway & Hostless via docker build",
        "Health check HTTP /health tiap 15 detik",
        "Env: DATABASE_URL, JWT_SECRET, STATIC_DIR",
        "Git untuk versi kode",
    ]),
]
sx = ML
for title_i, col, items in stacks:
    rect(s, sx, 1.7, 3.9, 4.9, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, sx, 1.7, 3.9, 0.6, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, sx, 1.86, 3.9, 0.35,
         [{"t": title_i, "size": 15, "bold": True, "color": WHITE}],
         align=PP_ALIGN.CENTER)
    iy = 2.5
    for it in items:
        rect(s, sx + 0.28, iy + 0.07, 0.11, 0.11, fill=col, shape=MSO_SHAPE.OVAL)
        text(s, sx + 0.52, iy, 3.15, 0.6,
             [{"t": it, "size": 11.5, "color": SLATE, "line_spacing": 1.2}])
        iy += 0.66
    sx += 4.02
footer(s, 7)

# ------------------------------------------------------- 8. endpoint
s = blank(prs)
header(s, "API", "Modul Endpoint & Aksesnya", GREEN)
data = [
    ["Modul", "Endpoint utama", "Metode", "Akses"],
    ["Auth", "/api/auth/register · /login · /auth/me", "POST · GET", "Publik / semua login"],
    ["Pengguna", "/api/users · /users/:id/role", "GET · POST · PUT", "Owner"],
    ["Dashboard", "/api/dashboard", "GET", "Semua role login"],
    ["Produk", "/api/products · /products/:id", "GET · POST · PUT · DELETE", "Baca: semua · Tulis: owner, admin"],
    ["Stok", "/api/stock/movements", "GET · POST", "Owner, admin"],
    ["Penjualan", "/api/sales · /sales/:id", "GET · POST", "Owner, admin"],
    ["Servis", "/api/services · /:id · /:id/parts · /:id/payments", "GET · POST · PUT", "Kerja: owner, admin, mekanik · Buat: owner, admin"],
    ["Pembelian", "/api/purchases · /purchases/:id", "GET · POST", "Owner, admin"],
]
make_table(s, ML, 1.7, 11.93, 3.9, data, [1.6, 4.5, 2.4, 3.43],
           header_fill=NAVY, size=10.5, header_size=11.5, row_h=0.42)

rect(s, ML, 5.85, 11.93, 0.9, fill=LIGHT, line=LINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, ML + 0.3, 6.02, 11.4, 0.6, [
    {"t": "Catatan: 26 endpoint di bawah /api  ·  semua daftar mendukung "
          "pagination ?page=1&limit=20 (maks 100) dan mengembalikan meta "
          "page/limit/total  ·  respons JSON konsisten dengan field "
          "data & error.", "size": 11.5, "color": SLATE,
     "line_spacing": 1.25},
])
footer(s, 8)

# ------------------------------------------------- 9. keamanan/role
s = blank(prs)
header(s, "Keamanan", "Kendali Akses Berbasis Role (RBAC)", RED)
role_data = [
    ["Modul", "Owner", "Admin", "Mekanik", "Staff"],
    ["Dashboard", "✓", "✓", "✓", "✓"],
    ["Inventaris (baca)", "✓", "✓", "✓", "✓"],
    ["Penjualan", "✓", "✓", "—", "—"],
    ["Servis", "✓", "✓", "✓", "—"],
    ["Pembelian", "✓", "✓", "—", "—"],
    ["Kelola produk & stok", "✓", "✓", "—", "—"],
    ["Kelola pengguna & role", "✓", "—", "—", "—"],
]
make_table(s, ML, 1.68, 6.4, 3.6, role_data, [2.4, 1.0, 1.0, 1.0, 1.0],
           header_fill=NAVY, size=11.5, header_size=11.5, row_h=0.44)

sec = [
    ("Autentikasi", "JWT HS256 · token 8 jam · secret wajib ≥ 32 karakter", BLUE),
    ("Password", "Tersimpan sebagai hash bcrypt, tidak pernah plaintext", GREEN),
    ("Rate limit", "Login 10 percobaan/menit · registrasi 5/menit per IP", AMBER),
    ("Validasi input", "Binding tag Gin + CHECK constraint di database", VIOLET),
    ("Integritas data", "Transaksi DB + kunci baris saat transaksi stok", RED),
]
sy = 1.68
for title_i, desc, col in sec:
    rect(s, 7.35, sy, 5.28, 0.9, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, 7.35, sy, 0.07, 0.9, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 7.6, sy + 0.15, 4.8, 0.3,
         [{"t": title_i, "size": 12.5, "bold": True, "color": NAVY}])
    text(s, 7.6, sy + 0.46, 4.8, 0.35,
         [{"t": desc, "size": 10.8, "color": SLATE}])
    sy += 1.0

text(s, ML, 5.6, 6.4, 0.9, [
    {"t": "User pertama yang mendaftar otomatis menjadi owner; user "
          "berikutnya berstatus staff dan diberi hak oleh owner. Menu di "
          "frontend otomatis menyesuaikan role (ProtectedRoute + RoleRoute).",
     "size": 11.5, "color": SLATE, "line_spacing": 1.3},
])
footer(s, 9)

# ---------------------------------------------------- 10. database
s = blank(prs)
header(s, "Database", "Skema Data (10 Tabel)", BLUE)
groups = [
    ("Pengguna", NAVY, ["users"]),
    ("Produk & Stok", BLUE, ["products", "stock_movements"]),
    ("Penjualan", GREEN, ["sales", "sale_items"]),
    ("Servis", AMBER, ["services", "service_parts", "service_payments"]),
    ("Pembelian", VIOLET, ["purchases", "purchase_items"]),
]
gx = ML
for title_i, col, tables in groups:
    rect(s, gx, 1.68, 2.3, 1.75, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, gx, 1.68, 2.3, 0.05, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, gx + 0.18, 1.85, 2.0, 0.3,
         [{"t": title_i, "size": 12.5, "bold": True, "color": col}])
    tb = []
    for t in tables:
        tb.append({"t": t, "size": 11.5, "bold": True, "color": NAVY,
                   "space_after": 4})
    text(s, gx + 0.18, 2.2, 2.0, 1.1, tb)
    gx += 2.4

text(s, ML, 3.7, 5.7, 0.35,
     [{"t": "Integritas & performa", "size": 14, "bold": True, "color": NAVY}])
bullets(s, ML, 4.15, 5.7, 2.5, [
    "Relasi FOREIGN KEY antar-tabel (hapus data tidak merusak)",
    "CHECK constraint: harga/qty tidak negatif, stok ≥ 0, status valid",
    "16 index pada kolom yang sering dicari & diurutkan",
    "Kolom generated (subtotal) agar konsisten",
    "10 file migration dijalankan otomatis saat server start",
], size=12, gap=11)

rect(s, 6.8, 3.7, 5.83, 2.9, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.1, 3.95, 5.3, 0.35,
     [{"t": "Alur data contoh: penjualan", "size": 13.5, "bold": True,
       "color": AMBER}])
flow = ["Validasi item → INSERT sales",
        "SELECT product ... FOR UPDATE",
        "INSERT sale_items",
        "UPDATE products (stok − qty)",
        "INSERT stock_movements (out)",
        "UPDATE total → COMMIT"]
fy = 4.4
for i, f in enumerate(flow):
    text(s, 7.1, fy, 5.3, 0.3,
         [{"t": f"{i + 1}.  {f}", "size": 11.5,
           "color": WHITE if i != 5 else GREEN,
           "bold": i == 5}])
    fy += 0.33
footer(s, 10)

# ---------------------------- 11. data demo, pengujian, deployment
s = blank(prs)
header(s, "Kesiapan demo", "Data Demo, Pengujian & Deployment", AMBER)
cols = [
    ("Data demo otomatis", GREEN, [
        "56 jenis spare part (SKU, kategori, harga, stok)",
        "5 akun berbeda role siap dipakai",
        "3 pembelian, 35 penjualan (54 item)",
        "15 pekerjaan servis + 7 pembayaran tunai",
        "84 pergerakan stok yang konsisten",
        "Tercetak untuk 7 hari terakhir — dashboard langsung terisi",
    ]),
    ("Pengujian & kualitas", BLUE, [
        "Unit test Go: pagination, validasi penjualan",
        "Unit test: kalkulasi stok & proses seed",
        "Validasi input dua lapis: binding + constraint DB",
        "Pagination & meta konsisten di semua daftar",
        "ESLint + tsc -b saat build frontend",
        "Mode seed aman: hanya saat tabel produk kosong",
    ]),
    ("Deployment", VIOLET, [
        "Docker multi-stage: build React → build Go → Alpine",
        "Hasil build frontend disajikan backend (satu origin)",
        "Railway / Hostless cukup dengan Dockerfile",
        "Health check /health tiap 15 detik",
        "Non-root user & timeout server",
        "Skalabel ke hosting mana pun",
    ]),
]
sx = ML
for title_i, col, items in cols:
    rect(s, sx, 1.7, 3.9, 4.9, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, sx, 1.7, 3.9, 0.06, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, sx + 0.25, 1.95, 3.4, 0.35,
         [{"t": title_i, "size": 14, "bold": True, "color": NAVY}])
    iy = 2.45
    for it in items:
        rect(s, sx + 0.28, iy + 0.07, 0.1, 0.1, fill=col, shape=MSO_SHAPE.OVAL)
        text(s, sx + 0.5, iy, 3.2, 0.65,
             [{"t": it, "size": 11, "color": SLATE, "line_spacing": 1.2}])
        iy += 0.68
    sx += 4.02
footer(s, 11)

# ---------------------------------------------------- 12. penutup
s = blank(prs)
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, 0, 0.22, SH, fill=BLUE)
text(s, 1.0, 0.9, 8.0, 0.4,
     [{"t": "KESIMPULAN", "size": 13, "bold": True, "color": AMBER}])
text(s, 1.0, 1.3, 11.0, 0.7,
     [{"t": "Satu aplikasi untuk seluruh operasional bengkel",
       "size": 30, "bold": True, "color": WHITE}])

concl = [
    "Seluruh proses bengkel — pembelian, stok, penjualan, servis, dan "
    "pembayaran — tercatat digital dalam satu basis data.",
    "Keamanan berlapis: bcrypt, JWT, rate limit, dan pembagian 4 role yang "
    "konsisten di API maupun tampilan.",
    "Siap dipresentasikan dan di-deploy: data demo otomatis, unit test, "
    "migrasi otomatis, dan satu image Docker.",
]
cy = 2.25
for i, c in enumerate(concl, 1):
    rect(s, 1.0, cy, 0.5, 0.5, fill=BLUE, shape=MSO_SHAPE.OVAL)
    text(s, 1.0, cy + 0.1, 0.5, 0.3,
         [{"t": str(i), "size": 15, "bold": True, "color": WHITE}],
         align=PP_ALIGN.CENTER)
    text(s, 1.7, cy + 0.02, 6.4, 0.9,
         [{"t": c, "size": 13.5, "color": RGBColor(0xD3, 0xDC, 0xEB),
           "line_spacing": 1.3}])
    cy += 1.05

rect(s, 8.5, 2.2, 3.9, 2.9, fill=RGBColor(0x16, 0x21, 0x3A), line=BLUE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.8, 2.45, 3.4, 0.35,
     [{"t": "Rekomendasi pengembangan", "size": 13, "bold": True,
       "color": AMBER}])
bullets(s, 8.8, 2.9, 3.3, 2.1, [
    "Cetak nota / struk PDF",
    "Pembayaran QRIS & e-wallet",
    "Laporan laba rugi per periode",
    "Notifikasi stok minimum",
    "Riwayat servis per kendaraan",
], size=11.5, color=RGBColor(0xC3, 0xCE, 0xE0), gap=8)

rect(s, 1.0, 5.55, 11.4, 0.9, fill=RGBColor(0x16, 0x21, 0x3A),
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.3, 5.72, 10.8, 0.6, [
    {"t": "Terima kasih  ·  Sesi demo langsung", "size": 17, "bold": True,
     "color": WHITE, "space_after": 3},
    {"t": "owner@bengkelku.id / owner123   ·   admin@bengkelku.id / admin123"
          "   ·   budi@bengkelku.id / mekanik123", "size": 11.5,
     "color": RGBColor(0x9A, 0xAB, 0xC6)},
])
text(s, 1.0, 6.75, 11.0, 0.3,
     [{"t": "WorkshopPro  ·  Tugas Akhir Pelatihan BLKPP  ·  Salma Auliana  ·  "
           "github.com/MangWifiId", "size": 10.5,
       "color": RGBColor(0x8A, 0x9A, 0xB5)}])

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Presentasi_WorkshopPro.pptx")
prs.save(out)
print("Tersimpan:", out)
print("Jumlah slide:", len(prs.slides._sldIdLst))
